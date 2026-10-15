"""0.7 automation blocks: text conditions, working hours, JSON APIs, Discord/Slack, moving processed files."""
from __future__ import annotations

import datetime as dt
import json
import pathlib
import re
import shutil
import threading
import time
import urllib.request

from ..ai.config import home
from ..core import Node, Param, PortType, register
from .ai_hub import _Feed, _Job
from .outputs import _rising

NUM, BOOL, TEXT, ANY = PortType.NUMBER, PortType.BOOL, PortType.TEXT, PortType.ANY
LOGIC, OUT, HUB = "4 · Mantık", "5 · Çıkış ve Otomasyon", "6 · AI Merkezi"

_TR = str.maketrans({"İ": "i", "I": "ı", "Ş": "ş", "Ğ": "ğ", "Ü": "ü", "Ö": "ö", "Ç": "ç"})


def fold(text: str) -> str:
    """Case-insensitive form that keeps Turkish letters right ("İSTANBUL" == "istanbul")."""
    return str(text).translate(_TR).lower()


@register
class TextMatch(Node):
    setup_params = ("case", "mode", "words")   # read only when the block starts: changing them restarts it
    title = "Metin Koşulu"
    category = LOGIC
    description = ("Metin bir kelimeyi içeriyor mu, onunla mı başlıyor, ya da bir kalıba (regex) uyuyor mu? "
                   "Virgülle birden çok kelime yazılabilir: 'fatura, ödeme, iade'. Büyük/küçük harf ve Türkçe "
                   "harfler doğru karşılaştırılır. E-postaları, mesajları konuya göre ayırmak için.")
    inputs = [("text", ANY)]
    outputs = [("match", BOOL), ("found", TEXT), ("not_match", BOOL)]
    params = [Param("mode", "choice", "içeriyor", "Koşul", choices=["içeriyor", "ile başlıyor", "eşit", "regex"]),
              Param("words", "str", "fatura, ödeme", "Kelimeler (virgülle) / kalıp"),
              Param("case", "bool", False, "Büyük/küçük harf duyarlı")]

    def setup(self):
        self.state["rx"] = None
        if self.v("mode") == "regex":
            try:
                self.state["rx"] = re.compile(self.v("words"), 0 if self.v("case") else re.IGNORECASE)
            except re.error as exc:
                raise ValueError(f"Kalıp hatalı: {exc}") from None

    def process(self, inputs):
        text = inputs["text"]
        if text is None:
            return {"match": False, "found": "", "not_match": True}
        text = str(text)
        mode, found = self.v("mode"), ""
        if mode == "regex":
            m = self.state["rx"].search(text) if self.state["rx"] else None
            found = (m.group(1) if m and m.groups() else m.group(0)) if m else ""
            hit = m is not None
        else:
            norm = (lambda s: s) if self.v("case") else fold
            hay = norm(text).strip()
            hit = False
            for w in (w.strip() for w in self.v("words").split(",")):
                if not w:
                    continue
                n = norm(w)
                if (mode == "içeriyor" and n in hay) or (mode == "ile başlıyor" and hay.startswith(n)) or \
                        (mode == "eşit" and hay == n):
                    hit, found = True, w
                    break
        return {"match": hit, "found": found or "", "not_match": not hit}


def _hm(text: str) -> dt.time:
    try:
        h, m = (int(x) for x in str(text).strip().replace(".", ":").split(":"))
        return dt.time(h, m)
    except (ValueError, TypeError):
        raise ValueError(f"Saat '{text}' anlaşılamadı, 09:00 gibi yazın") from None


DAYS = {"her gün": range(7), "hafta içi": range(5), "hafta sonu": range(5, 7)}


@register
class TimeWindow(Node):
    setup_params = ("end", "start")   # read only when the block starts: changing them restarts it
    title = "Saat Aralığında"
    category = LOGIC
    description = ("Şu an belirlenen saatler arasındaysa doğru verir: mesai saati, gece vardiyası (22:00-06:00 "
                   "gibi gece yarısını geçen aralıklar da olur). Bildirimleri sadece mesai saatinde göndermek için "
                   "Mantık Kapısı ile birleştirin.")
    outputs = [("active", BOOL), ("inactive", BOOL)]
    params = [Param("start", "str", "09:00", "Başlangıç"), Param("end", "str", "18:00", "Bitiş"),
              Param("days", "choice", "hafta içi", "Günler", choices=list(DAYS))]

    def now(self) -> dt.datetime:
        return dt.datetime.now()

    def setup(self):
        self.state.update(start=_hm(self.v("start")), end=_hm(self.v("end")))

    def process(self, inputs):
        now = self.now()
        s, e, t = self.state["start"], self.state["end"], now.time()
        if s <= e:
            on, day = s <= t < e, now.weekday()
        else:                     # overnight: 22:00-06:00 belongs to the day it started
            on = t >= s or t < e
            day = now.weekday() if t >= s else (now.weekday() - 1) % 7
        on = on and day in DAYS[self.v("days")]
        return {"active": on, "inactive": not on}


def json_path(data, path: str):
    """'data.items.0.price' → data["data"]["items"][0]["price"]; empty path = whole document."""
    cur = data
    for part in [p for p in re.split(r"[.\[\]]", path or "") if p]:
        if isinstance(cur, list):
            try:
                cur = cur[int(part)]
            except (ValueError, IndexError):
                raise KeyError(f"'{part}' liste içinde bulunamadı") from None
        elif isinstance(cur, dict):
            if part not in cur:
                raise KeyError(f"'{part}' bulunamadı (var olanlar: {', '.join(list(cur)[:10])})")
            cur = cur[part]
        else:
            raise KeyError(f"'{part}' için önce bir nesne/liste gerekir")
    return cur


@register
class ApiGet(_Job):
    title = "API'den Veri Al"
    category = HUB
    description = ("Bir web adresinden JSON veri okur ve içinden istediğiniz alanı verir: döviz kuru, hava durumu, "
                   "stok, kendi sunucunuz. Alan yolu örneği: 'rates.USD' veya 'items.0.price'. Aralık verilirse "
                   "düzenli yeniler.")
    inputs = [("url", TEXT), ("trigger", BOOL)]
    outputs = [("value", ANY), ("number", NUM), ("text", TEXT), ("busy", BOOL), ("done", BOOL)]
    params = [Param("url", "str", "https://api.frankfurter.app/latest?from=USD&to=TRY", "Adres (giriş bağlı değilse)"),
              Param("path", "str", "rates.TRY", "Alan yolu (boş = hepsi)"),
              Param("header", "str", "", "Ek başlık (Ad: değer, isteğe bağlı)"),
              Param("interval", "float", 0.0, "Yenileme aralığı (sn, 0 = sadece tetik/değişince)", 0, 86400)]
    key = "url"
    empty = {"value": None, "number": 0.0, "text": ""}

    def setup(self):
        super().setup()
        self.state["next_t"] = 0.0

    def should_fire(self, inputs):
        if inputs.get("url") is None:
            inputs["url"] = self.v("url")
        fire = super().should_fire(inputs)
        if self.v("interval") > 0 and time.monotonic() >= self.state["next_t"]:
            fire = True
        if fire:
            self.state["next_t"] = time.monotonic() + (self.v("interval") or 1e12)
        return fire

    def run(self, inputs):
        url = str(inputs.get("url") or self.v("url")).strip()
        if not url.startswith(("http://", "https://")):
            raise ValueError("Adres http:// veya https:// ile başlamalı")
        headers = {"User-Agent": "VisionFlow", "Accept": "application/json"}
        if ":" in self.v("header"):
            k, v = self.v("header").split(":", 1)
            headers[k.strip()] = v.strip()
        with urllib.request.urlopen(urllib.request.Request(url, headers=headers), timeout=15) as r:
            raw = r.read(5_000_000).decode("utf-8", "replace")
        try:
            data = json.loads(raw)
        except ValueError:
            raise ValueError("Cevap JSON değil") from None
        value = json_path(data, self.v("path"))
        try:
            number = float(value)
        except (TypeError, ValueError):
            number = 0.0
        text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False)
        return {"value": value, "number": number, "text": text}


@register
class ChatWebhook(Node):
    title = "Discord / Slack / Teams Gönder"
    category = OUT
    description = ("Tetik yükseldiğinde bir kanala mesaj yazar. Kanal ayarlarından 'Webhook' adresi alıp yapıştırın; "
                   "Discord, Slack ve Teams adresleri otomatik tanınır. {v} yerine gelen değer yazılır.")
    inputs = [("trigger", BOOL), ("value", ANY)]
    outputs = [("sent", NUM)]
    params = [Param("url", "str", "", "Webhook adresi"), Param("text", "code", "Olay: {v}", "Mesaj"),
              Param("cooldown", "float", 10.0, "En az bekleme (sn)", 0, 86400)]

    LIMITS = {"discord": 2000, "slack": 40000, "teams": 28000}

    def setup(self):
        self.state.update(last=-1e18, sent=0)

    @staticmethod
    def kind(url: str) -> str:
        u = url.lower()
        return "discord" if "discord" in u else "teams" if ("office.com" in u or "logic.azure" in u) else "slack"

    @classmethod
    def payload(cls, url: str, text: str) -> dict:
        kind = cls.kind(url)
        text = text[:cls.LIMITS[kind]]
        return {"content": text} if kind == "discord" else {"text": text}

    def process(self, inputs):
        if not _rising(self, inputs["trigger"]) or not self.v("url"):
            return {"sent": self.state["sent"]}
        now = time.monotonic()
        if now - self.state["last"] < self.v("cooldown"):
            return {"sent": self.state["sent"]}
        self.state["last"] = now
        v = inputs["value"]
        text = self.v("text").replace("{v}", "" if v is None else str(v))
        body = json.dumps(self.payload(self.v("url"), text)).encode()
        req = urllib.request.Request(self.v("url"), body, {"Content-Type": "application/json",
                                                           "User-Agent": "VisionFlow"})

        def send():
            try:
                urllib.request.urlopen(req, timeout=10).close()
                self.state["sent"] += 1
            except Exception as exc:  # noqa: BLE001
                print(f"[{self.title}] gönderilemedi: {exc}")
        threading.Thread(target=send, daemon=True).start()
        return {"sent": self.state["sent"]}


def free_name(folder: pathlib.Path, name: str) -> pathlib.Path:
    """folder/name, or folder/name_1, name_2... if taken, so nothing is ever overwritten."""
    target = folder / name
    stem, suffix, i = target.stem, target.suffix, 1
    while target.exists():
        target = folder / f"{stem}_{i}{suffix}"
        i += 1
    return target


@register
class MoveFile(Node):
    title = "Dosyayı Taşı / Kopyala"
    category = OUT
    description = ("Tetik yükseldiğinde dosyayı başka bir klasöre taşır veya kopyalar; aynı adda dosya varsa üzerine "
                   "yazmaz, '_1' ekler. Klasörü İzle ile işlenen dosyaları 'islendi' klasörüne almak için.")
    inputs = [("path", TEXT), ("trigger", BOOL)]
    outputs = [("new_path", TEXT), ("moved", NUM)]
    params = [Param("folder", "folder", "islendi", "Hedef klasör"),
              Param("mode", "choice", "taşı", "İşlem", choices=["taşı", "kopyala"]),
              Param("dated", "bool", False, "Tarihli alt klasör (2026-10-12)")]

    def setup(self):
        self.state.update(new_path="", moved=0)

    def process(self, inputs):
        fire = _rising(self, inputs["trigger"]) if inputs["trigger"] is not None else \
            inputs["path"] not in (None, "", self.state.get("last_path")) and \
            time.monotonic() >= self.state.get("retry_at", 0)
        src = inputs["path"]
        if fire and src:
            src = pathlib.Path(src)
            if not src.is_file():
                self.state["last_path"] = inputs["path"]
                raise FileNotFoundError(f"Dosya yok: {src}")
            folder = pathlib.Path(self.v("folder"))
            if self.v("dated"):
                folder = folder / dt.date.today().isoformat()
            folder.mkdir(parents=True, exist_ok=True)
            if folder.resolve() == src.parent.resolve():
                self.state["last_path"] = inputs["path"]
                return {"new_path": str(src), "moved": self.state["moved"]}
            target = free_name(folder, src.name)
            try:
                (shutil.move if self.v("mode") == "taşı" else shutil.copy2)(str(src), str(target))
            except OSError as exc:
                # e.g. the file is still being written / open in another program: across drives shutil.move
                # copies first and then fails to delete the original, leaving two files. Undo the copy so
                # the next try starts clean, and try again a little later.
                if src.exists() and target.exists():
                    try:
                        target.unlink()
                    except OSError:
                        pass
                self.state["retry_at"] = time.monotonic() + 2
                raise RuntimeError(f"{src.name} taşınamadı (başka bir program kullanıyor olabilir), tekrar "
                                   f"denenecek: {exc}") from None
            self.state["last_path"] = inputs["path"]
            self.state.update(new_path=str(target), moved=self.state["moved"] + 1)
        return {"new_path": self.state["new_path"], "moved": self.state["moved"]}


# --- 0.8 -----------------------------------------------------------------------------------------------

def parse_feed(raw: bytes) -> list[dict]:
    """RSS 2.0 / RSS 1.0 / Atom → [{'id', 'title', 'link', 'summary', 'date'}], newest first as published."""
    import xml.etree.ElementTree as ET
    from .ai_hub import html_to_text
    try:
        root = ET.fromstring(raw)
    except ET.ParseError as exc:
        raise ValueError(f"Haber akışı okunamadı (RSS/Atom değil): {exc}") from None

    def local(tag):
        return tag.rsplit("}", 1)[-1].lower()

    def child(el, *names):
        for c in el:
            if local(c.tag) in names:
                return c
        return None

    def text(el, *names):
        c = child(el, *names)
        return (c.text or "").strip() if c is not None else ""

    items = []
    for el in root.iter():
        kind = local(el.tag)
        if kind not in ("item", "entry"):
            continue
        link = text(el, "link")
        if not link:                                    # Atom: <link href="..." rel="alternate"/>
            for c in el:
                if local(c.tag) == "link" and c.get("rel", "alternate") == "alternate" and c.get("href"):
                    link = c.get("href")
                    break
        summary = text(el, "description", "summary", "content", "encoded")
        title = html_to_text(text(el, "title"))[1]
        items.append({"id": text(el, "guid", "id") or link or title, "title": title, "link": link,
                      "summary": html_to_text(summary)[1][:4000],
                      "date": text(el, "pubdate", "published", "updated", "date")})
    return items


@register
class RssFeed(_Feed):
    setup_params = ("url", "existing")
    title = "Haber Akışı (RSS)"
    category = HUB
    description = ("Haber sitelerinin, blogların, YouTube kanallarının RSS/Atom akışını izler; yeni çıkan her "
                   "yazıyı sırayla verir. Birden çok adres virgülle yazılabilir. Kelime yazılırsa sadece onları "
                   "içeren haberler gelir. 'Yapay Zekaya Sor' ile özetletip Telegram'a gönderebilirsiniz.")
    inputs = [("next", BOOL)]
    outputs = [("title", TEXT), ("link", TEXT), ("summary", TEXT), ("text", TEXT), ("new", BOOL), ("waiting", NUM)]
    params = [Param("url", "str", "https://www.aa.com.tr/tr/rss/default?cat=guncel", "Akış adresi (virgülle)"),
              Param("words", "str", "", "Sadece bu kelimeler (virgülle, boş = hepsi)"),
              Param("interval", "float", 600.0, "Kontrol aralığı (sn)", 30, 86400),
              Param("existing", "bool", False, "Başlarken mevcut haberleri de ver")]

    def setup(self):
        super().setup()
        self.state["first"] = not self.v("existing")

    def _seen_file(self):
        import hashlib
        return home() / "haberler" / f"{hashlib.sha1(self.v('url').encode()).hexdigest()[:12]}.json"

    def fetch(self):
        urls = [u.strip() for u in self.v("url").split(",") if u.strip()]
        if not urls:
            raise ValueError("Akış adresi boş")
        items = []
        for u in urls:
            if not u.startswith(("http://", "https://")):
                raise ValueError(f"Adres http:// ile başlamalı: {u}")
            req = urllib.request.Request(u, headers={"User-Agent": "Mozilla/5.0 (VisionFlow)"})
            with urllib.request.urlopen(req, timeout=20) as r:
                items += list(reversed(parse_feed(r.read(10_000_000))))     # oldest first
        words = [fold(w.strip()) for w in self.v("words").split(",") if w.strip()]
        if words:
            items = [i for i in items if any(w in fold(i["title"] + " " + i["summary"]) for w in words)]
        if self.state["first"] and not self.state["seen"]:
            # first run: what is already published counts as read, only news from now on is handed out
            self.state["first"] = False
            for i in items:
                self._remember(self.item_key(i))
            self._save_seen()
            return []
        self.state["first"] = False
        return items

    def item_key(self, item):
        return item["id"]

    def empty_outputs(self):
        return {"title": "", "link": "", "summary": "", "text": ""}

    def outputs_for(self, item):
        return {"title": item["title"], "link": item["link"], "summary": item["summary"],
                "text": f"{item['title']}\n{item['summary']}\n{item['link']}".strip()}


@register
class Delay(Node):
    title = "Gecikme"
    category = LOGIC
    description = ("Gelen tetiği ve değeri N saniye sonra iletir (her tetik ayrı ayrı, kaybolmaz). Ör. parça "
                   "algılandıktan 1,5 sn sonra iticiyi çalıştırmak.")
    inputs = [("trigger", BOOL), ("value", ANY)]
    outputs = [("out", BOOL), ("value", ANY), ("waiting", NUM)]
    params = [Param("seconds", "float", 1.0, "Gecikme (sn)", 0, 86400, 0.1)]

    def setup(self):
        self.state.update(queue=[], value=None)

    def process(self, inputs):
        now = time.monotonic()
        q = self.state["queue"]
        if _rising(self, inputs["trigger"]) and inputs["trigger"] is not None:
            q.append((now + self.v("seconds"), inputs["value"]))
        out = False
        if q and q[0][0] <= now:          # one per frame, so each delayed trigger is its own pulse
            _, self.state["value"] = q.pop(0)
            out = True
        return {"out": out, "value": self.state["value"], "waiting": len(q)}


@register
class DailyCounter(Node):
    setup_params = ("name", "reset_time")   # read only when the block starts: changing them restarts it
    title = "Günlük Sayaç"
    category = LOGIC
    description = ("Tetikleri sayar, her gün belirlenen saatte sıfırlanır (vardiya, günlük üretim). Sayı "
                   "kaydedilir: program kapanıp açılınca kaldığı yerden devam eder. Dünün sayısını da verir.")
    inputs = [("trigger", BOOL), ("reset", BOOL)]
    outputs = [("today", NUM), ("yesterday", NUM), ("total", NUM)]
    params = [Param("name", "str", "uretim", "Sayaç adı"), Param("reset_time", "str", "00:00", "Sıfırlama saati")]

    def now(self) -> dt.datetime:
        return dt.datetime.now()

    def _file(self):
        return home() / "sayaclar" / f"{re.sub(r'[^\w-]+', '_', self.v('name')) or 'sayac'}.json"

    def _day(self, t: dt.datetime) -> str:
        """The counting day a moment belongs to: with reset at 06:00, 03:00 still counts for yesterday."""
        r = self.state["reset"]
        return (t.date() - dt.timedelta(days=1) if t.time() < r else t.date()).isoformat()

    def setup(self):
        self.state.update(reset=_hm(self.v("reset_time")), prev=False)
        try:
            data = json.loads(self._file().read_text(encoding="utf-8"))
        except (OSError, ValueError):
            data = {}
        self.state["data"] = {"day": data.get("day", ""), "today": int(data.get("today", 0)),
                              "yesterday": int(data.get("yesterday", 0)), "total": int(data.get("total", 0))}

    def _save(self):
        f = self._file()
        f.parent.mkdir(parents=True, exist_ok=True)
        tmp = f.with_suffix(".tmp")
        tmp.write_text(json.dumps(self.state["data"]), encoding="utf-8")
        tmp.replace(f)

    def process(self, inputs):
        d, day = self.state["data"], self._day(self.now())
        changed = False
        if d["day"] != day:
            prev_day = (dt.date.fromisoformat(day) - dt.timedelta(days=1)).isoformat()
            d.update(yesterday=d["today"] if d["day"] == prev_day else 0, today=0, day=day)
            changed = True
        if inputs["reset"]:
            d["today"], changed = 0, True
        now = bool(inputs["trigger"])
        if now and not self.state["prev"]:
            d["today"] += 1
            d["total"] += 1
            changed = True
        self.state["prev"] = now
        if changed:
            self._save()
        return {"today": d["today"], "yesterday": d["yesterday"], "total": d["total"]}


@register
class ExcelLogger(Node):
    title = "Excel'e Kaydet"
    category = OUT
    description = ("Değerleri tarih-saatle birlikte gerçek bir Excel dosyasına (.xlsx) satır satır yazar. Dosya "
                   "Excel'de açıkken yazılamazsa satırlar bekletilir, dosya kapanınca hepsi yazılır, kayıp olmaz.")
    inputs = [("a", ANY), ("b", ANY), ("c", ANY), ("trigger", BOOL)]
    outputs = [("rows", NUM), ("waiting", NUM)]
    params = [Param("path", "str", "kayit.xlsx", "Dosya"), Param("sheet", "str", "Kayıtlar", "Sayfa"),
              Param("names", "str", "a,b,c", "Sütun adları")]

    SAVE_EVERY = 1.0        # seconds: rows arriving faster are written together (opening .xlsx is slow)
    RETRY_EVERY = 3.0       # seconds between attempts while the file is locked or broken

    def setup(self):
        self.state["last"] = None
        import importlib.util
        if importlib.util.find_spec("openpyxl") is None:
            raise RuntimeError("Excel için 'openpyxl' paketi gerekli: pip install openpyxl")
        self.state.update(pending=[], rows=0, next_try=0.0, error=None)

    def teardown(self):
        if self.state.get("pending"):
            try:
                self._flush()
            except Exception:  # noqa: BLE001 - stopping the flow must not fail because Excel holds the file
                pass

    @staticmethod
    def _cell(v):
        if v is None or isinstance(v, (int, float, str, bool)):
            return v
        if hasattr(v, "item"):                     # numpy scalar
            return v.item()
        return str(v)

    def _flush(self):
        import openpyxl
        path = pathlib.Path(self.v("path"))
        sheet = (self.v("sheet") or "Kayıtlar")[:31]
        names = [n.strip() for n in self.v("names").split(",")][:3]
        if path.exists():
            try:
                wb = openpyxl.load_workbook(path)
            except PermissionError:
                raise
            except Exception as exc:  # noqa: BLE001 - zip/xml errors of a damaged or non-Excel file
                raise RuntimeError(f"{path.name} bozuk ya da Excel dosyası değil ({type(exc).__name__}); "
                                   f"başka bir dosya adı seçin") from None
            ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet)
        else:
            path.parent.mkdir(parents=True, exist_ok=True)
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet
        if ws.max_row == 1 and all(c.value is None for c in ws[1]):     # new sheet: header in row 1
            for col, title in enumerate(["Zaman", *names], 1):
                ws.cell(1, col, title)
            ws.column_dimensions["A"].width = 20
            r = 2
        else:
            r = ws.max_row + 1
        for row in self.state["pending"]:
            for col, value in enumerate(row, 1):
                cell = ws.cell(r, col, value)
                if isinstance(value, str) and value.startswith("="):
                    cell.data_type = "s"           # text that looks like a formula stays text, never runs
            r += 1
        wb.save(path)
        self.state["rows"] += len(self.state["pending"])
        self.state["pending"] = []

    def process(self, inputs):
        if inputs["trigger"] is not None and not _rising(self, inputs["trigger"]):
            values = None
        else:
            values = [self._cell(inputs[k]) for k in "abc"]
            if inputs["trigger"] is None and (all(v is None for v in values) or values == self.state.get("last")):
                values = None                      # no trigger: a row only when the values change
        if values is not None:
            self.state["last"] = values
            self.state["pending"].append([dt.datetime.now().replace(microsecond=0), *values])
        now = time.monotonic()
        if self.state["pending"] and now >= self.state["next_try"]:
            try:
                self._flush()
                self.state["error"], self.state["next_try"] = None, now + self.SAVE_EVERY
            except PermissionError:
                self.state["next_try"] = now + self.RETRY_EVERY
                self.state["error"] = "Dosya başka bir programda (Excel) açık; kapanınca hepsi yazılır"
            except RuntimeError as exc:
                self.state["next_try"] = now + self.RETRY_EVERY
                self.state["error"] = str(exc)
        if self.state["error"] and self.state["pending"]:
            raise RuntimeError(f"{len(self.state['pending'])} satır bekletiliyor: {self.state['error']}")
        return {"rows": self.state["rows"], "waiting": len(self.state["pending"])}
