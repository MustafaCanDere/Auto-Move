"""PowerPoint presentation generation block: converts WhatsApp, Email, or AI input into PPTX presentations."""
from __future__ import annotations

import datetime as dt
import json
import os
import pathlib
import re
import tempfile
import time

import cv2
import numpy as np

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

from ..core import Node, Param, PortType, register
from ..core import node as _node
from .processing import _bgr


def _rising(node: Node, trigger, key: str = "prev_trigger") -> bool:
    if trigger is None:
        once_key = f"{key}_once"
        if not node.state.get(once_key, False):
            node.state[once_key] = True
            return True
        return False
    now = bool(trigger)
    fired = now and not node.state.get(key, False)
    node.state[key] = now
    return fired


CAT_OUT = "5 · Çıkış ve Otomasyon"

IMG = PortType.IMAGE
NUM = PortType.NUMBER
BOOL = PortType.BOOL
TEXT = PortType.TEXT
ANY = PortType.ANY


# Theme Palette Definitions (RGB)
THEMES = {
    "Modern Gece (Koyu Şık)": {
        "bg": RGBColor(15, 23, 42),         # Slate 900
        "card": RGBColor(30, 41, 59),       # Slate 800
        "primary": RGBColor(248, 250, 252), # Slate 50
        "secondary": RGBColor(148, 163, 184), # Slate 400
        "accent": RGBColor(6, 182, 212),    # Cyan 500
        "accent_dim": RGBColor(14, 116, 144),
        "tbl_hdr": RGBColor(14, 116, 144),
        "tbl_row": RGBColor(30, 41, 59),
        "tbl_alt": RGBColor(40, 53, 73),
    },
    "Kurumsal Mavi": {
        "bg": RGBColor(248, 250, 252),      # Light
        "card": RGBColor(255, 255, 255),
        "primary": RGBColor(15, 23, 42),    # Slate 900
        "secondary": RGBColor(71, 85, 105), # Slate 600
        "accent": RGBColor(30, 64, 175),    # Blue 800
        "accent_dim": RGBColor(59, 130, 246),
        "tbl_hdr": RGBColor(30, 64, 175),
        "tbl_row": RGBColor(241, 245, 249),
        "tbl_alt": RGBColor(255, 255, 255),
    },
    "Endüstriyel Turuncu": {
        "bg": RGBColor(24, 24, 27),         # Zinc 900
        "card": RGBColor(39, 39, 42),       # Zinc 800
        "primary": RGBColor(250, 250, 250),
        "secondary": RGBColor(161, 161, 170),
        "accent": RGBColor(249, 115, 22),   # Orange 500
        "accent_dim": RGBColor(194, 65, 12),
        "tbl_hdr": RGBColor(194, 65, 12),
        "tbl_row": RGBColor(39, 39, 42),
        "tbl_alt": RGBColor(50, 50, 54),
    },
    "Minimal Açık": {
        "bg": RGBColor(255, 255, 255),
        "card": RGBColor(249, 250, 251),
        "primary": RGBColor(17, 24, 39),
        "secondary": RGBColor(75, 85, 99),
        "accent": RGBColor(37, 99, 235),
        "accent_dim": RGBColor(96, 165, 250),
        "tbl_hdr": RGBColor(37, 99, 235),
        "tbl_row": RGBColor(243, 244, 246),
        "tbl_alt": RGBColor(255, 255, 255),
    },
}


def _parse_content_into_slides(raw_text: str, default_title: str) -> list[dict]:
    """Parse unstructured message, Markdown, or bullet lists into slide specifications."""
    slides = []
    lines = [line.strip() for line in (raw_text or "").split("\n") if line.strip()]

    if not lines:
        return [{
            "title": default_title or "Özet ve Durum",
            "subtitle": "Otomatik Oluşturulan Rapor",
            "bullets": ["Sistemden gelen veri veya mesaj bulunamadı."],
        }]

    # Check if text contains Markdown headers '# '
    has_md = any(line.startswith("# ") for line in lines)
    if has_md:
        current_slide = None
        for line in lines:
            if line.startswith("# "):
                if current_slide:
                    slides.append(current_slide)
                current_slide = {
                    "title": line[2:].strip(),
                    "subtitle": "",
                    "bullets": [],
                }
            elif line.startswith("## ") and current_slide:
                current_slide["subtitle"] = line[3:].strip()
            elif current_slide:
                bullet = re.sub(r"^[-*•\d.]+\s*", "", line)
                if bullet:
                    current_slide["bullets"].append(bullet)
        if current_slide:
            slides.append(current_slide)
    else:
        # Unstructured text (e.g. WhatsApp message or Email body)
        bullets = []
        for line in lines:
            # Check sentence endings or line items
            items = [item.strip() for item in re.split(r"(?<=[.!?])\s+", line) if item.strip()]
            for item in items:
                clean_item = re.sub(r"^[-*•\d.]+\s*", "", item)
                if len(clean_item) > 3:
                    bullets.append(clean_item)

        # Chunk bullets into slides of up to 4 items each
        chunk_size = 4
        for i in range(0, max(1, len(bullets)), chunk_size):
            chunk = bullets[i:i + chunk_size]
            idx = (i // chunk_size) + 1
            stitle = default_title if idx == 1 else f"{default_title} (Bölüm {idx})"
            slides.append({
                "title": stitle,
                "subtitle": f"Detay Özeti · Parça {idx}",
                "bullets": chunk if chunk else ["Bilgi kaydı mevcut değil."],
            })

    return slides


@register
class PresentationGenerate(Node):
    title = "Sunum Hazırla (PowerPoint)"
    category = CAT_OUT
    description = (
        "WhatsApp, e-posta veya yapay zekadan gelen metin ve tablolardan modern Microsoft PowerPoint (.pptx) "
        "sunumu hazırlar. Başlık slaytı, içerik slaytları, madde imleri, tablolar ve denetim görsellerini otomatik dizer."
    )
    inputs = [
        ("trigger", BOOL),
        ("title", TEXT),
        ("subtitle", TEXT),
        ("content", ANY),
        ("image", IMG),
        ("data", ANY),
    ]
    outputs = [
        ("path", TEXT),
        ("done", BOOL),
        ("slide_count", NUM),
    ]
    params = [
        Param("theme", "choice", "Modern Gece (Koyu Şık)", "Sunum Teması", choices=list(THEMES.keys())),
        Param("output_folder", "folder", "sunumlar", "Kayıt Klasörü"),
        Param("filename", "str", "sunum_{zaman}.pptx", "Dosya Adı ({zaman} otomatik eklenir)"),
        Param("author", "str", "VisionFlow AI Endüstriyel Çözümler", "Yazar / Kurum"),
        Param("aspect_ratio", "choice", "16:9 (Geniş Ekran)", "En/Boy Oranı",
              choices=["16:9 (Geniş Ekran)", "4:3 (Klasik Standart)"]),
    ]

    def setup(self):
        self.state.update(path="", done=False, slide_count=0, prev_trig=False)

    def process(self, inputs):
        st = self.state
        st.setdefault("path", "")
        st.setdefault("done", False)
        st.setdefault("slide_count", 0)

        if not _rising(self, inputs.get("trigger"), "prev_trig"):
            return {"path": st["path"], "done": False, "slide_count": st["slide_count"]}

        if not PPTX_AVAILABLE:
            raise RuntimeError("python-pptx kütüphanesi kurulu değil. Lütfen 'pip install python-pptx' çalıştırın.")

        theme_name = self.v("theme")
        colors = THEMES.get(theme_name, THEMES["Modern Gece (Koyu Şık)"])

        prs = Presentation()

        # Set slide dimensions
        if self.v("aspect_ratio").startswith("16:9"):
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)
        else:
            prs.slide_width = Inches(10)
            prs.slide_height = Inches(7.5)

        width = prs.slide_width
        height = prs.slide_height

        blank_slide_layout = prs.slide_layouts[6]  # completely blank layout

        def apply_bg(slide):
            bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, width, height)
            bg_shape.fill.solid()
            bg_shape.fill.fore_color.rgb = colors["bg"]
            bg_shape.line.fill.background()
            return bg_shape

        main_title = str(inputs.get("title") or "Üretim ve Kalite Değerlendirme Sunumu").strip()
        sub_title = str(inputs.get("subtitle") or "Otomasyon ve Görsel Denetim Raporu").strip()
        author_name = str(self.v("author") or "VisionFlow").strip()
        now_str = dt.datetime.now().strftime("%d.%m.%Y %H:%M")

        # ----------------------------------------------------------------------------------
        # SLIDE 1: Title Slide
        # ----------------------------------------------------------------------------------
        s1 = prs.slides.add_slide(blank_slide_layout)
        apply_bg(s1)

        # Decorative Top Accent Bar
        bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(1.2), Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = colors["accent"]
        bar.line.fill.background()

        # Title Card Box
        tb = s1.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.5), Inches(3.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = main_title
        p1.font.size = Pt(40)
        p1.font.bold = True
        p1.font.color.rgb = colors["primary"]

        if sub_title:
            p2 = tf.add_paragraph()
            p2.text = sub_title
            p2.font.size = Pt(22)
            p2.font.color.rgb = colors["accent"]
            p2.space_before = Pt(14)

        # Meta Card Bottom
        meta_box = s1.shapes.add_textbox(Inches(0.8), height - Inches(1.5), Inches(11.5), Inches(0.8))
        mtf = meta_box.text_frame
        mp = mtf.paragraphs[0]
        mp.text = f"📅 {now_str}   |   🏢 {author_name}   |   ⚡ VisionFlow AI Studio"
        mp.font.size = Pt(13)
        mp.font.color.rgb = colors["secondary"]

        # ----------------------------------------------------------------------------------
        # Handle Image if provided
        # ----------------------------------------------------------------------------------
        temp_img_file = None
        img_input = inputs.get("image")
        if img_input is not None:
            if isinstance(img_input, str):
                p = pathlib.Path(_node.resolve_path(img_input))
                if p.is_file():
                    temp_img_file = str(p)
            elif isinstance(img_input, np.ndarray):
                try:
                    bgr = _bgr(img_input)
                    tmp = tempfile.NamedTemporaryFile(prefix="vf_pptx_", suffix=".png", delete=False)
                    tmp.close()
                    cv2.imwrite(tmp.name, bgr)
                    temp_img_file = tmp.name
                except Exception as exc:
                    print(f"[Sunum Hazırla] Görsel kaydedilemedi: {exc}")

        # ----------------------------------------------------------------------------------
        # SLIDE 2+: Content Slides (parsed from WhatsApp/Email/AI text)
        # ----------------------------------------------------------------------------------
        raw_content = inputs.get("content")
        content_str = str(raw_content) if raw_content is not None else ""
        parsed_slides = _parse_content_into_slides(content_str, main_title)

        for s_idx, s_info in enumerate(parsed_slides):
            s = prs.slides.add_slide(blank_slide_layout)
            apply_bg(s)

            # Slide Header Bar
            hdr_shape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.6), Inches(0.2), Inches(0.6))
            hdr_shape.fill.solid()
            hdr_shape.fill.fore_color.rgb = colors["accent"]
            hdr_shape.line.fill.background()

            title_box = s.shapes.add_textbox(Inches(1.2), Inches(0.55), Inches(11.0), Inches(0.8))
            tt_frame = title_box.text_frame
            tt_p = tt_frame.paragraphs[0]
            tt_p.text = s_info.get("title", "Genel Durum")
            tt_p.font.size = Pt(26)
            tt_p.font.bold = True
            tt_p.font.color.rgb = colors["primary"]

            sub = s_info.get("subtitle", "")
            if sub:
                sub_p = tt_frame.add_paragraph()
                sub_p.text = sub
                sub_p.font.size = Pt(14)
                sub_p.font.color.rgb = colors["secondary"]

            # Layout: If image is available and this is the first content slide, place side by side
            has_img_on_this_slide = (temp_img_file and s_idx == 0)
            text_width = Inches(6.5) if has_img_on_this_slide else Inches(11.5)

            # Content Text Box / Bullets
            content_box = s.shapes.add_textbox(Inches(0.8), Inches(1.8), text_width, Inches(4.8))
            c_tf = content_box.text_frame
            c_tf.word_wrap = True

            bullets = s_info.get("bullets", [])
            for b_idx, bullet in enumerate(bullets):
                bp = c_tf.paragraphs[0] if b_idx == 0 else c_tf.add_paragraph()
                bp.text = f"•  {bullet}"
                bp.font.size = Pt(18)
                bp.font.color.rgb = colors["primary"]
                bp.space_before = Pt(12)
                bp.space_after = Pt(4)

            # Insert Image on side if present
            if has_img_on_this_slide:
                try:
                    img_left = Inches(7.7)
                    img_top = Inches(1.8)
                    img_width = Inches(4.8)
                    s.shapes.add_picture(temp_img_file, img_left, img_top, width=img_width)
                except Exception as img_exc:
                    print(f"[Sunum Hazırla] Slayta görsel eklenirken hata: {img_exc}")

            # Slide Footer
            ftr = s.shapes.add_textbox(Inches(0.8), height - Inches(0.8), Inches(11.5), Inches(0.5))
            ftp = ftr.text_frame.paragraphs[0]
            ftp.text = f"{author_name}   ·   {now_str}   ·   Slayt {s_idx + 2}"
            ftp.font.size = Pt(10)
            ftp.font.color.rgb = colors["secondary"]

        # ----------------------------------------------------------------------------------
        # SLIDE (OPTIONAL): Data Table Slide (if structured data provided)
        # ----------------------------------------------------------------------------------
        data_in = inputs.get("data")
        table_rows = []
        if isinstance(data_in, list) and data_in and isinstance(data_in[0], dict):
            table_rows = data_in[:12]  # limit to top 12 rows
        elif isinstance(data_in, str) and data_in.strip().startswith("["):
            try:
                parsed_json = json.loads(data_in)
                if isinstance(parsed_json, list) and parsed_json and isinstance(parsed_json[0], dict):
                    table_rows = parsed_json[:12]
            except Exception:
                pass

        if table_rows:
            ts = prs.slides.add_slide(blank_slide_layout)
            apply_bg(ts)

            # Header
            t_hdr = ts.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.6), Inches(0.2), Inches(0.6))
            t_hdr.fill.solid()
            t_hdr.fill.fore_color.rgb = colors["accent"]
            t_hdr.line.fill.background()

            t_tbox = ts.shapes.add_textbox(Inches(1.2), Inches(0.55), Inches(11.0), Inches(0.8))
            t_p = t_tbox.text_frame.paragraphs[0]
            t_p.text = "Detaylı Üretim & Kalite Veri Tablosu"
            t_p.font.size = Pt(26)
            t_p.font.bold = True
            t_p.font.color.rgb = colors["primary"]

            headers = list(table_rows[0].keys())[:7]
            n_rows = len(table_rows) + 1
            n_cols = len(headers)

            t_shape = ts.shapes.add_table(n_rows, n_cols, Inches(0.8), Inches(1.8), Inches(11.7), Inches(0.4 * n_rows))
            table = t_shape.table

            # Format Header Row
            for c_idx, h_text in enumerate(headers):
                cell = table.cell(0, c_idx)
                cell.fill.solid()
                cell.fill.fore_color.rgb = colors["tbl_hdr"]
                cp = cell.text_frame.paragraphs[0]
                cp.text = str(h_text).upper()
                cp.font.size = Pt(13)
                cp.font.bold = True
                cp.font.color.rgb = RGBColor(255, 255, 255)
                cp.alignment = PP_ALIGN.CENTER

            # Format Data Rows
            for r_idx, row_dict in enumerate(table_rows):
                for c_idx, h_text in enumerate(headers):
                    cell = table.cell(r_idx + 1, c_idx)
                    cell.fill.solid()
                    row_bg = colors["tbl_row"] if r_idx % 2 == 0 else colors["tbl_alt"]
                    cell.fill.fore_color.rgb = row_bg
                    cp = cell.text_frame.paragraphs[0]
                    cp.text = str(row_dict.get(h_text, ""))
                    cp.font.size = Pt(12)
                    cp.font.color.rgb = colors["primary"]
                    cp.alignment = PP_ALIGN.CENTER

        # Save presentation
        folder = self.v("output_folder") or "sunumlar"
        folder_path = pathlib.Path(_node.resolve_path(folder))
        folder_path.mkdir(parents=True, exist_ok=True)

        stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        fname_template = str(self.v("filename") or "sunum_{zaman}.pptx")
        fname = fname_template.replace("{zaman}", stamp)
        if not fname.endswith(".pptx"):
            fname += ".pptx"

        out_path = folder_path / fname
        prs.save(str(out_path))

        st["path"] = str(out_path.resolve())
        st["done"] = True
        st["slide_count"] = len(prs.slides)

        print(f"[Sunum Hazırla] PowerPoint oluşturuldu: {st['path']} ({st['slide_count']} slayt)")
        return {"path": st["path"], "done": True, "slide_count": st["slide_count"]}
